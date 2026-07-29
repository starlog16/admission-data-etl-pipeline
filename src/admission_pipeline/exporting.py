from __future__ import annotations

import json
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from .domain import ParsedRecord


def _headers(records: list[ParsedRecord]) -> list[str]:
    ordered: list[str] = []
    for record in records:
        for key in record.values:
            if not key.startswith("_") and key not in ordered:
                ordered.append(key)
    return ordered


def _excel_scalar(value: object) -> object:
    if isinstance(value, (dict, list, tuple, set)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def _field_evidence_text(record: ParsedRecord) -> str:
    evidence = record.values.get("_field_evidence", {})
    if not isinstance(evidence, dict):
        return ""
    parts: list[str] = []
    for field, details in evidence.items():
        if not isinstance(details, dict):
            continue
        raw = details.get("raw_value", "")
        final = details.get("final_value", raw)
        cell = details.get("cell_value", "")
        status = details.get("status", "")
        confidence = details.get("raw_confidence", "")
        cell_confidence = details.get("cell_confidence", "")
        parts.append(
            f"{field}: raw={raw}({confidence}), cell={cell}({cell_confidence}), "
            f"final={final}, status={status}"
        )
    return "；".join(parts)


def _write_sheet(
    sheet: Any,
    records: list[ParsedRecord],
    *,
    include_issues: bool,
    fallback_headers: list[str] | None = None,
) -> None:
    headers = _headers(records) or list(fallback_headers or [])
    if include_issues:
        headers.extend(["关键字段证据", "审计说明", "待复核原因"])
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for record in records:
        values = []
        for header in headers:
            if header == "待复核原因":
                continue
            if header == "关键字段证据":
                values.append(_field_evidence_text(record))
            elif header == "审计说明":
                values.append("；".join(record.warnings))
            else:
                values.append(_excel_scalar(record.values.get(header, "")))
        if include_issues:
            values.append("；".join(record.issues))
        sheet.append(values)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        maximum = max(
            [len(str(header))]
            + [len(str(sheet.cell(row=row, column=index).value or "")) for row in range(2, sheet.max_row + 1)]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(maximum + 2, 10), 42)


def _write_audit_sheet(sheet: Any, records: list[ParsedRecord]) -> None:
    headers = _headers(records) + ["关键字段证据", "审计说明"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for record in records:
        sheet.append(
            [
                *[
                    _field_evidence_text(record)
                    if header == "关键字段证据"
                    else _excel_scalar(record.values.get(header, ""))
                    for header in headers[:-1]
                ],
                "；".join(record.warnings),
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _write_rows(sheet: Any, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        maximum = max(
            [len(str(header))]
            + [len(str(sheet.cell(row=row, column=index).value or "")) for row in range(2, sheet.max_row + 1)]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(maximum + 2, 10), 42)


def _add_review_cases(workbook: Workbook, records: list[ParsedRecord]) -> None:
    review = [record for record in records if record.needs_review]
    if not review:
        return
    grouped: dict[tuple[str, ...], list[ParsedRecord]] = {}
    for record in review:
        source = str(record.values.get("来源文件", ""))
        issues = "；".join(sorted(record.issues))
        institution = str(
            record.values.get("院校名称", "") or record.values.get("院校代码", "")
        )
        group = str(
            record.values.get("专业组代码", "")
            or record.values.get("院校专业组代码", "")
        )
        parent_level = any(
            marker in issues for marker in ("院校代码", "院校名称", "专业组", "计划数合计")
        )
        key = (
            source,
            institution if parent_level else str(record.values.get("来源页", "")),
            group if parent_level else str(record.values.get("来源行", "")),
            issues,
        )
        grouped.setdefault(key, []).append(record)

    rows: list[list[object]] = []
    for case_number, items in enumerate(grouped.values(), start=1):
        first = items[0]
        pages = sorted({str(item.values.get("来源页", "")) for item in items})
        major_codes = sorted(
            {
                str(item.values.get("专业代码", ""))
                for item in items
                if str(item.values.get("专业代码", ""))
            }
        )
        rows.append(
            [
                f"R{case_number:05d}",
                len(items),
                "；".join(first.issues),
                first.values.get("来源文件", ""),
                ",".join(pages),
                first.values.get("院校代码", ""),
                first.values.get("院校名称", ""),
                first.values.get("专业组代码", "")
                or first.values.get("院校专业组代码", ""),
                ",".join(major_codes[:20]),
                _field_evidence_text(first),
            ]
        )
    _write_rows(
        workbook.create_sheet("review_cases"),
        [
            "复核案例ID",
            "影响记录数",
            "问题类型",
            "来源文件",
            "来源页",
            "院校代码",
            "院校名称",
            "专业组代码",
            "专业代码样例",
            "关键字段证据",
        ],
        rows,
    )


def _add_hierarchy_summaries(workbook: Workbook, records: list[ParsedRecord]) -> None:
    hierarchy = [
        record for record in records if str(record.values.get("院校专业组代码", "")).strip()
    ]
    if not hierarchy:
        return

    groups: dict[tuple[str, str, str, str], list[ParsedRecord]] = {}
    institutions: dict[tuple[str, str], list[ParsedRecord]] = {}
    for record in hierarchy:
        source = str(record.values.get("来源文件", ""))
        full_name = str(record.values.get("院校全称", "") or record.values.get("院校名称", ""))
        institution_code = str(record.values.get("院校代码", "")).strip()
        group_code = str(record.values.get("院校专业组代码", ""))
        course = str(record.values.get("科类", ""))
        subject = str(record.values.get("再选科目", "") or record.values.get("科目要求", ""))
        groups.setdefault((source, course, subject, group_code), []).append(record)
        institutions.setdefault((source, institution_code or full_name), []).append(record)

    group_rows: list[list[object]] = []
    for (source, course, subject, group_code), items in groups.items():
        first = items[0].values
        full_name = str(first.get("院校全称", "") or first.get("院校名称", ""))
        expected = str(first.get("专业组计划数", "")).strip()
        plans = [str(item.values.get("计划数", "")).strip() for item in items]
        actual: object = sum(int(value) for value in plans) if plans and all(value.isdigit() for value in plans) else ""
        result = "一致" if expected.isdigit() and actual == int(expected) else "待复核"
        if expected.isdigit() and isinstance(actual, int) and actual != int(expected):
            result = "不一致"
        group_rows.append([
            full_name, group_code, course, subject,
            first.get("院校专业组名称", ""), expected, actual,
            len(items), min(int(item.values.get("来源页", 1)) for item in items),
            result, source,
        ])

    institution_rows: list[list[object]] = []
    for (source, _institution_key), items in institutions.items():
        first = items[0].values
        institution_code = str(first.get("院校代码", "")).strip()
        names = [
            str(item.values.get("院校全称", "") or item.values.get("院校名称", "")).strip()
            for item in items
        ]
        full_name = Counter(name for name in names if name).most_common(1)[0][0]
        expected = str(first.get("院校计划总数", "")).strip()
        plans = [str(item.values.get("计划数", "")).strip() for item in items]
        actual: object = sum(int(value) for value in plans) if plans and all(value.isdigit() for value in plans) else ""
        result = "一致" if expected.isdigit() and actual == int(expected) else "待复核"
        if expected.isdigit() and isinstance(actual, int) and actual != int(expected):
            result = "不一致"
        institution_rows.append([
            institution_code, full_name, expected, actual,
            len({str(item.values.get("院校专业组代码", "")) for item in items}),
            len(items), min(int(item.values.get("来源页", 1)) for item in items),
            result, source,
        ])

    _write_rows(
        workbook.create_sheet("专业组汇总"),
        ["院校全称", "院校专业组代码", "科类", "再选科目", "院校专业组名称", "专业组计划数", "专业计划数合计", "专业数量", "来源页码", "校验结果", "来源文件"],
        group_rows,
    )
    _write_rows(
        workbook.create_sheet("院校汇总"),
        ["院校代码", "院校全称", "院校计划总数", "专业计划数合计", "专业组数量", "专业数量", "起始页码", "校验结果", "来源文件"],
        institution_rows,
    )


def export_workbook(
    output_path: Path,
    records: list[ParsedRecord],
    summary: dict[str, object],
    code_fields: list[str],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "data"
    good = [record for record in records if not record.needs_review]
    review = [record for record in records if record.needs_review]
    all_headers = _headers(records)
    _write_sheet(
        data_sheet,
        good,
        include_issues=False,
        fallback_headers=all_headers,
    )
    if review:
        review_sheet = workbook.create_sheet("review")
        _write_sheet(
            review_sheet,
            review,
            include_issues=True,
            fallback_headers=all_headers,
        )
        _add_review_cases(workbook, review)
    audit = [record for record in records if record.warnings and not record.needs_review]
    if audit:
        _write_audit_sheet(workbook.create_sheet("audit"), audit)

    _add_hierarchy_summaries(workbook, records)

    metadata = workbook.create_sheet("metadata")
    metadata.append(["项目", "值"])
    metadata.append(["生成时间", datetime.now().isoformat(timespec="seconds")])
    for key, value in summary.items():
        metadata.append([key, _excel_scalar(value)])

    for sheet in workbook.worksheets:
        header_map = {cell.value: cell.column for cell in sheet[1]}
        for field in code_fields:
            column = header_map.get(field)
            if column:
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row=row, column=column).number_format = "@"
    workbook.save(output_path)
    check = load_workbook(output_path, read_only=True)
    try:
        if not check.sheetnames:
            raise RuntimeError(f"Excel 重新打开校验失败：{output_path}")
    finally:
        check.close()
