from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import numpy as np
import pymupdf
from openpyxl import load_workbook
from PIL import Image
from rapidocr import RapidOCR

from .config import ProjectConfig
from .domain import ExtractedPage, InputFile, Token


class ExtractionEngine:
    """统一提取图片、文本 PDF 和扫描 PDF，并缓存昂贵的 OCR 结果。"""

    def __init__(self, config: ProjectConfig) -> None:
        self.config = config
        self._ocr: RapidOCR | None = None
        self.last_cache_hit = False

    @property
    def ocr(self) -> RapidOCR:
        if self._ocr is None:
            self._ocr = RapidOCR(
                params={
                    "Global.max_side_len": int(
                        self.config.extraction["ocr_global_max_side_len"]
                    )
                }
            )
        return self._ocr

    def extract(
        self,
        item: InputFile,
        *,
        use_cache: bool = True,
        force_ocr: bool = False,
    ) -> list[ExtractedPage]:
        cache_path = self._cache_path(item, force_ocr)
        if use_cache and cache_path.exists():
            self.last_cache_hit = True
            payload = json.loads(cache_path.read_text(encoding="utf-8"))
            pages = [ExtractedPage.from_dict(page) for page in payload["pages"]]
            if item.kind == "pdf" and any(
                page.method == "pdf_text" and not page.tables for page in pages
            ):
                self._attach_pdf_tables(item, pages)
                self._write_cache(cache_path, item, pages)
            return pages

        self.last_cache_hit = False

        if item.kind == "image":
            pages = [self._extract_image(item)]
        elif item.kind == "pdf":
            pages = self._extract_pdf(item, force_ocr=force_ocr)
        elif item.kind == "spreadsheet":
            pages = self._extract_spreadsheet(item)
        else:
            raise ValueError(f"未知输入类型：{item.kind}")

        self._write_cache(cache_path, item, pages)
        return pages

    def _write_cache(
        self, cache_path: Path, item: InputFile, pages: list[ExtractedPage]
    ) -> None:
        cache_path.parent.mkdir(parents=True, exist_ok=True)
        cache_path.write_text(
            json.dumps(
                {
                    "cache_version": self.config.extraction["cache_version"],
                    "source": item.to_dict(),
                    "pages": [page.to_dict() for page in pages],
                },
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )

    def _attach_pdf_tables(
        self, item: InputFile, pages: list[ExtractedPage]
    ) -> None:
        with pymupdf.open(item.path) as document:
            for extracted, pdf_page in zip(pages, document, strict=False):
                if extracted.method == "pdf_text" and not extracted.tables:
                    extracted.tables = _tables_from_pdf_page(pdf_page)

    def _extract_spreadsheet(self, item: InputFile) -> list[ExtractedPage]:
        """Convert workbook cells to the same positioned-token model used by PDFs."""
        workbook = load_workbook(item.path, read_only=False, data_only=True)
        pages: list[ExtractedPage] = []
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                max_column = max(1, sheet.max_column)
                max_row = max(1, sheet.max_row)
                tokens: list[Token] = []
                text_rows: list[str] = []
                table_rows: list[list[str]] = []
                for row_index, row in enumerate(sheet.iter_rows(), start=1):
                    texts: list[str] = []
                    table_row: list[str] = []
                    for column_index, cell in enumerate(row, start=1):
                        text = _spreadsheet_cell_text(cell.value, cell.number_format)
                        table_row.append(text)
                        if not text:
                            continue
                        texts.append(text)
                        tokens.append(
                            Token(
                                text=text,
                                score=1.0,
                                x1=float(column_index - 1),
                                y1=float(row_index - 1),
                                x2=float(column_index),
                                y2=float(row_index),
                            )
                        )
                    if texts:
                        text_rows.append("\t".join(texts))
                    table_rows.append(table_row)
                pages.append(
                    ExtractedPage(
                        source_relative_path=item.relative_path.as_posix(),
                        source_sha256=item.sha256,
                        page_number=sheet_index,
                        method="spreadsheet_cells",
                        width=float(max_column),
                        height=float(max_row),
                        tokens=tokens,
                        plain_text="\n".join(text_rows),
                        sheet_name=sheet.title,
                        tables=[table_rows],
                    )
                )
        finally:
            workbook.close()
        return pages

    def _cache_path(self, item: InputFile, force_ocr: bool) -> Path:
        identity = (
            f"{item.relative_path.as_posix()}|{item.sha256}|"
            f"{self.config.extraction['cache_version']}|{force_ocr}"
        )
        key = hashlib.sha256(identity.encode("utf-8")).hexdigest()
        return self.config.path("temp_root") / "cache" / "extractions" / f"{key}.json"

    def _extract_image(self, item: InputFile) -> ExtractedPage:
        with Image.open(item.path) as image:
            width, height = image.size
        return self._run_ocr(
            item.path,
            relative_path=item.relative_path.as_posix(),
            source_sha256=item.sha256,
            page_number=1,
            width=width,
            height=height,
            method="image_ocr",
        )

    def _extract_pdf(self, item: InputFile, *, force_ocr: bool) -> list[ExtractedPage]:
        pages: list[ExtractedPage] = []
        min_chars = int(self.config.extraction["pdf_text_min_chars_per_page"])
        dpi = int(self.config.extraction["pdf_render_dpi"])
        with pymupdf.open(item.path) as document:
            for index, page in enumerate(document):
                tokens = _tokens_from_pdf_page(page)
                plain_text = "\n".join(token.text for token in tokens)
                if not force_ocr and len("".join(plain_text.split())) >= min_chars:
                    rect = page.rect
                    pages.append(
                        ExtractedPage(
                            source_relative_path=item.relative_path.as_posix(),
                            source_sha256=item.sha256,
                            page_number=index + 1,
                            method="pdf_text",
                            width=float(rect.width),
                            height=float(rect.height),
                            tokens=tokens,
                            plain_text=plain_text,
                            tables=_tables_from_pdf_page(page),
                        )
                    )
                    continue

                scale = dpi / 72.0
                pixmap = page.get_pixmap(matrix=pymupdf.Matrix(scale, scale), alpha=False)
                image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
                pages.append(
                    self._run_ocr(
                        np.asarray(image),
                        relative_path=item.relative_path.as_posix(),
                        source_sha256=item.sha256,
                        page_number=index + 1,
                        width=pixmap.width,
                        height=pixmap.height,
                        method="pdf_page_ocr",
                    )
                )
        return pages

    def _run_ocr(
        self,
        source: str | Path | np.ndarray,
        *,
        relative_path: str,
        source_sha256: str,
        page_number: int,
        width: int,
        height: int,
        method: str,
    ) -> ExtractedPage:
        result = self.ocr(source)
        tokens: list[Token] = []
        if result.boxes is not None and result.txts is not None and result.scores is not None:
            minimum_score = float(self.config.extraction["ocr_min_score"])
            for box, text, score in zip(result.boxes, result.txts, result.scores, strict=False):
                score_value = float(score)
                if not str(text).strip() or score_value < minimum_score:
                    continue
                points = np.asarray(box, dtype=float)
                tokens.append(
                    Token(
                        text=str(text).strip(),
                        score=score_value,
                        x1=float(points[:, 0].min()),
                        y1=float(points[:, 1].min()),
                        x2=float(points[:, 0].max()),
                        y2=float(points[:, 1].max()),
                    )
                )
        tokens.sort(key=lambda token: (token.y1, token.x1))
        return ExtractedPage(
            source_relative_path=relative_path,
            source_sha256=source_sha256,
            page_number=page_number,
            method=method,
            width=float(width),
            height=float(height),
            tokens=tokens,
            plain_text="\n".join(token.text for token in tokens),
        )


def _tokens_from_pdf_page(page: pymupdf.Page) -> list[Token]:
    tokens: list[Token] = []
    page_dict: dict[str, Any] = page.get_text("dict", sort=True)
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = str(span.get("text", "")).strip()
                bbox = span.get("bbox")
                if not text or not bbox:
                    continue
                tokens.append(
                    Token(
                        text=text,
                        score=1.0,
                        x1=float(bbox[0]),
                        y1=float(bbox[1]),
                        x2=float(bbox[2]),
                        y2=float(bbox[3]),
                    )
                )
    return sorted(tokens, key=lambda token: (token.y1, token.x1))


def _tables_from_pdf_page(page: pymupdf.Page) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    for table in page.find_tables().tables:
        rows: list[list[str]] = []
        for row in table.extract():
            rows.append(["" if value is None else str(value) for value in row])
        if rows:
            tables.append(rows)
    return tables


def _spreadsheet_cell_text(value: object, number_format: str) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, int):
        format_text = str(number_format or "").strip()
        if format_text and set(format_text) == {"0"}:
            return str(value).zfill(len(format_text))
        return str(value)
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else str(value)
    return str(value).strip()
