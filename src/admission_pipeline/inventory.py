from __future__ import annotations

import hashlib
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from PIL import Image

from .config import ProjectConfig
from .domain import InputFile


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def resolve_input_target(config: ProjectConfig, user_path: str) -> Path:
    raw_root = config.path("input_root")
    candidate = Path(user_path)
    if not candidate.is_absolute():
        candidate = raw_root / candidate
    candidate = candidate.resolve()
    if not candidate.is_relative_to(raw_root):
        raise ValueError(f"输入必须位于 {raw_root} 之下")
    if not candidate.exists():
        raise FileNotFoundError(candidate)
    return candidate


def _build_item(
    path: Path,
    raw_root: Path,
    image_exts: set[str],
    pdf_exts: set[str],
    spreadsheet_exts: set[str],
) -> InputFile:
    suffix = path.suffix.lower()
    common = dict(
        path=path,
        relative_path=path.relative_to(raw_root),
        size_bytes=path.stat().st_size,
        sha256=sha256_file(path),
    )
    if suffix in image_exts:
        with Image.open(path) as image:
            width, height = image.size
        return InputFile(kind="image", width=width, height=height, **common)
    if suffix in pdf_exts:
        with pymupdf.open(path) as document:
            return InputFile(kind="pdf", page_count=document.page_count, **common)
    if suffix in spreadsheet_exts:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return InputFile(kind="spreadsheet", page_count=len(workbook.sheetnames), **common)
        finally:
            workbook.close()
    raise ValueError(f"不支持的文件类型：{path}")


def inventory(config: ProjectConfig, target: Path) -> list[InputFile]:
    image_exts = {value.lower() for value in config.extraction["image_extensions"]}
    pdf_exts = {value.lower() for value in config.extraction["pdf_extensions"]}
    spreadsheet_exts = {
        value.lower() for value in config.extraction.get("spreadsheet_extensions", [])
    }
    supported = image_exts | pdf_exts | spreadsheet_exts
    paths = [target] if target.is_file() else [p for p in target.rglob("*") if p.is_file()]
    return [
        _build_item(
            path,
            config.path("input_root"),
            image_exts,
            pdf_exts,
            spreadsheet_exts,
        )
        for path in sorted(paths)
        if path.suffix.lower() in supported
    ]
