from pathlib import Path

from admission_pipeline.config import ProjectConfig
from admission_pipeline.cli import _compact_daily_result, build_parser
from admission_pipeline.exporting import _excel_scalar
from admission_pipeline.exporting import export_workbook
from admission_pipeline.domain import ParsedRecord
from openpyxl import load_workbook
from admission_pipeline.pipeline import (
    _add_coverage_release_gate,
    _aggregate_output_path,
    _archive_existing_output,
    _unverified_profile_names,
)


def test_directory_target_creates_one_mirrored_workbook(tmp_path: Path) -> None:
    raw_root = tmp_path / "input" / "raw"
    target = raw_root / "anhui" / "2026" / "19安徽"
    target.mkdir(parents=True)
    config = ProjectConfig(
        root=tmp_path,
        raw={
            "paths": {
                "input_root": "input/raw",
                "excel_root": "output/excel",
            }
        },
    )
    path = _aggregate_output_path(config, target)
    assert path == tmp_path / "output" / "excel" / "anhui" / "2026" / "19安徽" / "19安徽.xlsx"


def test_check_command_is_available_without_excel_output_flag() -> None:
    args = build_parser().parse_args(["check", "--input", "anhui/2026/demo"])
    assert args.command == "check"
    assert not hasattr(args, "overwrite")


def test_daily_command_has_no_overwrite_or_no_cache_escape_hatch() -> None:
    args = build_parser().parse_args(["daily", "--input", "anhui/2026/demo"])
    assert args.command == "daily"
    assert not hasattr(args, "overwrite")
    assert not hasattr(args, "no_cache")


def test_metadata_mapping_is_serialized_for_excel() -> None:
    assert _excel_scalar({"b": 2, "a": 1}) == '{"a": 1, "b": 2}'


def test_daily_output_is_compact() -> None:
    compact = _compact_daily_result(
        {
            "status": "completed",
            "target": "demo",
            "file_count": 2,
            "files": [{"cache_hit": True}, {"cache_hit": False}],
            "summary": {
                "record_count": 10,
                "review_count": 1,
                "warning_count": 2,
                "duplicate_count": 0,
                "profiles": {"demo": 10},
            },
            "outputs": [{"path": "result.xlsx"}],
            "report_path": "report.json",
        }
    )
    assert compact["cache_hit_count"] == 1
    assert compact["new_extraction_count"] == 1
    assert compact["warning_count"] == 2
    assert "files" not in compact


def test_warning_only_records_stay_in_data_and_are_repeated_in_audit(tmp_path: Path) -> None:
    output = tmp_path / "audit.xlsx"
    record = ParsedRecord(
        values={"院校代码": "0001", "投档人数": "0"},
        warnings=["零投档"],
    )
    export_workbook(output, [record], {"record_count": 1}, ["院校代码"])
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook["data"].max_row == 2
        assert workbook["audit"].max_row == 2
        assert "review" not in workbook.sheetnames
    finally:
        workbook.close()


def test_existing_output_is_archived_without_overwrite(tmp_path: Path) -> None:
    output = tmp_path / "output" / "excel" / "anhui" / "2026" / "2026.xlsx"
    output.parent.mkdir(parents=True)
    output.write_text("old", encoding="utf-8")
    config = ProjectConfig(
        root=tmp_path,
        raw={
            "paths": {
                "excel_root": "output/excel",
                "archive_root": "output/archive",
            }
        },
    )
    archived = _archive_existing_output(config, output)
    assert archived.read_text(encoding="utf-8") == "old"
    assert not output.exists()


def test_unknown_profile_blocks_final_export() -> None:
    assert _unverified_profile_names(
        {
            "known": {"name": "known", "review_all": False},
            "generic": {"name": "generic-review", "review_all": True},
        }
    ) == ["generic-review"]


def test_unobservable_parser_blocks_formal_release() -> None:
    summary = {"release_safe": True, "release_blockers": []}
    reports = [
        {
            "parse_diagnostics": {
                "coverage_observable": False,
                "source_units_detected": 0,
                "records_emitted": 10,
                "recovered_units": 0,
                "unparsed_units": 0,
            }
        }
    ]

    _add_coverage_release_gate(summary, reports, 1)

    assert not summary["release_safe"]
    assert summary["coverage"]["unobservable_file_count"] == 1
    assert "尚不能独立证明完整率" in summary["release_blockers"][0]


def test_hierarchical_output_adds_group_and_institution_summaries(tmp_path: Path) -> None:
    output = tmp_path / "result.xlsx"
    records = [
        ParsedRecord(values={
            "院校全称": "测试大学（河北）", "院校计划总数": "2",
            "院校专业组代码": "000101", "科目要求": "不限",
            "院校专业组名称": "测试大学01组", "专业组计划数": "2",
            "专业代码": "01", "计划数": "2", "来源页": "1", "来源文件": "demo.pdf",
        })
    ]
    export_workbook(output, records, {"record_count": 1}, ["院校专业组代码", "专业代码"])
    workbook = load_workbook(output, read_only=True)
    try:
        assert "专业组汇总" in workbook.sheetnames
        assert "院校汇总" in workbook.sheetnames
        group_headers = {
            cell.value: cell.column for cell in workbook["专业组汇总"][1]
        }
        institution_headers = {
            cell.value: cell.column for cell in workbook["院校汇总"][1]
        }
        assert workbook["专业组汇总"].cell(
            row=2, column=group_headers["校验结果"]
        ).value == "一致"
        assert workbook["院校汇总"].cell(
            row=2, column=institution_headers["校验结果"]
        ).value == "一致"
    finally:
        workbook.close()


def test_export_serializes_visible_container_values(tmp_path: Path) -> None:
    output = tmp_path / "container.xlsx"
    record = ParsedRecord(
        values={"院校代码": "0001", "辅助信息": []},
        issues=[],
    )
    export_workbook(output, [record], {"record_count": 1}, ["院校代码"])
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook["data"].cell(row=2, column=2).value == "[]"
    finally:
        workbook.close()


def test_review_output_contains_grouped_cases_and_field_evidence(tmp_path: Path) -> None:
    output = tmp_path / "review.xlsx"
    records = [
        ParsedRecord(
            values={
                "来源文件": "fujian/demo.pdf",
                "来源页": 2,
                "院校代码": "1906",
                "院校名称": "福州理工学院",
                "专业代码": code,
                "_field_evidence": {
                    "院校代码": {
                        "raw_value": "1906",
                        "raw_confidence": 0.68,
                        "cell_value": "",
                        "cell_confidence": 0.0,
                        "final_value": "1906",
                        "status": "unresolved",
                    }
                },
            },
            issues=["院校代码网格单元格双路复识未达可信阈值"],
        )
        for code in ("021", "023")
    ]

    export_workbook(output, records, {"record_count": 2}, ["院校代码", "专业代码"])
    workbook = load_workbook(output, read_only=True)
    try:
        assert workbook["data"].max_row == 1
        assert workbook["review"].max_row == 3
        assert workbook["review_cases"].max_row == 2
        headers = {cell.value: cell.column for cell in workbook["review"][1]}
        evidence = workbook["review"].cell(
            row=2, column=headers["关键字段证据"]
        ).value
        assert "raw=1906" in evidence
        assert "status=unresolved" in evidence
        case_headers = {
            cell.value: cell.column for cell in workbook["review_cases"][1]
        }
        assert workbook["review_cases"].cell(
            row=2, column=case_headers["影响记录数"]
        ).value == 2
    finally:
        workbook.close()
